import zipfile
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from archivador.__main__ import main
from archivador.descarga import datos_xml, descargar, escribir_indice, escribir_zip, ruta_de, trae_cufe
from archivador.token import clase, leer_solo, leer_token, nit_empresa

EMPRESA = "901682408"


def _token(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Reporte de documentos"])
    ws.append([])
    ws.append(["Tipo de documento", "CUFE/CUDE", "Folio", "Prefijo", "Fecha Emisión", "NIT Emisor", "Nombre Emisor", "NIT Receptor", "Nombre Receptor", "Total"])
    ws.append(["Factura electrónica", "CUFE-A", "10", "FV", "15-09-2026", "900111222", "FERRETERIA", EMPRESA, "MI EMPRESA", "1.190.000,00"])
    ws.append(["Factura electrónica", "CUFE-B", "164", "FE", datetime(2026, 9, 20), EMPRESA, "MI EMPRESA", "800100200", "CLIENTE", 357000])
    ws.append(["Application response", "EVENTO-1", "1", "", "16-09-2026", EMPRESA, "MI EMPRESA", "900111222", "FERRETERIA", 0])
    ws.append(["Nómina Individual", "CUNE-C", "70", "NI", "30/09/2026", EMPRESA, "MI EMPRESA", "1020304050", "ANA", 2500000])
    ws.append(["Documento soporte", "CUDS-D", "27", "DSE", "2026-09-16", EMPRESA, "MI EMPRESA", "71000000", "JARDINERO", 300000])
    ws.append(["Factura electrónica", "cufe-a", "10", "FV", "15-09-2026", "900111222", "FERRETERIA", EMPRESA, "MI EMPRESA", 1190000])   # repetido
    ruta = tmp_path / "token.xlsx"
    wb.save(ruta)
    return ruta


def _xml(uuid: str) -> bytes:
    return (f'<Invoice xmlns="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2" xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2">'
            f"<cbc:UUID>{uuid}</cbc:UUID></Invoice>").encode()


class DianFalsa:
    def __init__(self, respuestas):
        self.respuestas = respuestas
        self.pedidos = []

    def get_xml(self, cufe):
        self.pedidos.append(cufe)
        r = self.respuestas[cufe]
        if isinstance(r, Exception):
            raise r
        return r


def test_lee_el_token_sin_eventos_ni_repetidos(tmp_path):
    docs = leer_token(_token(tmp_path))
    assert [d.cufe for d in docs] == ["CUFE-A", "CUFE-B", "CUNE-C", "CUDS-D"]
    a = docs[0]
    assert (a.numero, a.fecha, a.total, a.nit_emisor) == ("FV10", date(2026, 9, 15), Decimal("1190000.00"), "900111222")
    assert docs[1].fecha == date(2026, 9, 20) and docs[2].fecha == date(2026, 9, 30)
    assert nit_empresa(docs) == EMPRESA
    assert [clase(d, EMPRESA) for d in docs] == ["compras", "ventas", "nomina", "documento_soporte"]


def test_valida_el_cufe_del_xml_incluido_attached_document():
    assert trae_cufe(_xml("CUFE-A"), "cufe-a")
    assert not trae_cufe(_xml("OTRO"), "CUFE-A")
    assert not trae_cufe(b"no es xml", "CUFE-A")
    adjunto = (b'<AttachedDocument xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2"><cbc:Description>'
               b"&lt;Invoice&gt;&lt;cbc:UUID&gt;CUFE-A&lt;/cbc:UUID&gt;&lt;/Invoice&gt;</cbc:Description></AttachedDocument>")
    assert trae_cufe(adjunto, "CUFE-A")


def test_descarga_en_lote_organiza_salta_lo_que_existe_y_reporta_errores(tmp_path):
    docs = leer_token(_token(tmp_path))
    destino = tmp_path / "Facturacion"
    ya = ruta_de(destino, docs[3], "documento_soporte")
    ya.parent.mkdir(parents=True); ya.write_bytes(_xml("CUDS-D"))
    fuente = DianFalsa({"CUFE-A": _xml("CUFE-A"), "CUFE-B": _xml("OTRO"), "CUNE-C": RuntimeError("código 66")})
    esperas = []
    res = descargar(docs, fuente, destino, EMPRESA, pausa=0.5, dormir=esperas.append)
    assert [r.estado for r in res] == ["descargado", "error", "error", "ya_estaba"]
    assert fuente.pedidos == ["CUFE-A", "CUFE-B", "CUNE-C"]           # el que ya estaba no se pide
    assert esperas == [0.5, 0.5]                                         # pausa entre descargas
    assert res[0].ruta == destino / "2026" / "09" / "compras" / "FV10_900111222.xml"
    assert "no trae ese CUFE" in res[1].detalle and "66" in res[2].detalle

    indice = load_workbook(escribir_indice(res, destino / "indice.xlsx"))
    assert indice.sheetnames == ["Indice", "No descargados"]
    assert indice["Indice"].max_row == 5 and indice["No descargados"].max_row == 3
    with zipfile.ZipFile(escribir_zip(res, tmp_path / "lote.zip", destino)) as z:
        assert sorted(z.namelist()) == ["2026/09/compras/FV10_900111222.xml", "2026/09/documento_soporte/DSE27_71000000.xml"]

    # Segunda corrida: lo descargado ya no se vuelve a pedir.
    fuente2 = DianFalsa({"CUFE-B": _xml("CUFE-B"), "CUNE-C": _xml("CUNE-C")})
    res2 = descargar(docs, fuente2, destino, EMPRESA, pausa=0)
    assert fuente2.pedidos == ["CUFE-B", "CUNE-C"] and [r.estado for r in res2] == ["ya_estaba", "descargado", "descargado", "ya_estaba"]


def test_cli_con_la_lista_de_faltantes(tmp_path, capsys):
    token = _token(tmp_path)
    faltantes = tmp_path / "faltantes.csv"
    faltantes.write_text("cufe;clase;documento\ncufe-b;Venta;FE164\n", encoding="utf-8")
    assert leer_solo(faltantes) == {"cufe-b"}
    fuente = DianFalsa({"CUFE-B": _xml("CUFE-B")})
    codigo = main(["--desde-excel", str(token), "--solo", str(faltantes), "--destino", str(tmp_path / "out"), "--zip", str(tmp_path / "lote.zip"), "--pausa", "0"], cliente=fuente)
    salida = capsys.readouterr().out
    assert codigo == 0 and fuente.pedidos == ["CUFE-B"]
    assert "Descargados 1" in salida and (tmp_path / "out" / "indice.xlsx").exists() and (tmp_path / "lote.zip").exists()


def test_cli_sin_token_valido(tmp_path, capsys):
    malo = tmp_path / "malo.csv"
    malo.write_text("a,b\n1,2\n", encoding="utf-8")
    assert main(["--desde-excel", str(malo), "--destino", str(tmp_path)], cliente=DianFalsa({})) == 1
    assert "CUFE/CUDE" in capsys.readouterr().err


UBL = ('xmlns="urn:oasis:names:specification:ubl:schema:xsd:{raiz}-2" '
       'xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2" '
       'xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2"')


def _parte(rol, nit, nombre):
    return (f"<cac:{rol}><cac:Party><cac:PartyTaxScheme><cbc:RegistrationName>{nombre}</cbc:RegistrationName>"
            f'<cbc:CompanyID schemeID="5">{nit}</cbc:CompanyID></cac:PartyTaxScheme></cac:Party></cac:{rol}>')


def _factura(uuid, numero, fecha_, emisor, receptor, total, tipo="01"):
    return (f'<Invoice {UBL.format(raiz="Invoice")}><cbc:ID>{numero}</cbc:ID><cbc:UUID>{uuid}</cbc:UUID><cbc:IssueDate>{fecha_}</cbc:IssueDate>'
            f"<cbc:InvoiceTypeCode>{tipo}</cbc:InvoiceTypeCode>{_parte('AccountingSupplierParty', emisor, 'EMISOR SAS')}"
            f"{_parte('AccountingCustomerParty', receptor, 'RECEPTOR SAS')}"
            f'<cac:LegalMonetaryTotal><cbc:PayableAmount currencyID="COP">{total}</cbc:PayableAmount></cac:LegalMonetaryTotal></Invoice>')


def _adjunto(factura):
    return (f'<AttachedDocument {UBL.format(raiz="AttachedDocument")}><cbc:ParentDocumentID>X</cbc:ParentDocumentID>'
            f"<cac:Attachment><cac:ExternalReference><cbc:Description><![CDATA[{factura}]]></cbc:Description></cac:ExternalReference></cac:Attachment>"
            f"</AttachedDocument>").encode()


def test_datos_del_xml_factura_adjunto_soporte_y_nomina():
    d = datos_xml(_factura("u", "FE164", "2026-09-20", "901682408", "800100200", "357000.00").encode())
    assert (d["tipo"], d["folio"], d["fecha"], d["nit_emisor"], d["nit_receptor"], d["total"]) == (
        "Factura electrónica", "FE164", date(2026, 9, 20), "901682408", "800100200", Decimal("357000.00"))
    assert datos_xml(_adjunto(_factura("u", "FV10", "2026-09-15", "900111222", EMPRESA, "1")))["folio"] == "FV10"
    assert datos_xml(_factura("u", "DSE27", "2026-09-16", EMPRESA, "71000000", "1", tipo="05").encode())["tipo"] == "Documento soporte"
    nomina = (b'<NominaIndividual xmlns="dian:gov:co:facturaelectronica:NominaIndividual"><NumeroSecuenciaXML Numero="NI70"/>'
              b'<InformacionGeneral FechaGen="2026-09-30"/><Empleador NIT="901682408" RazonSocial="MI EMPRESA"/>'
              b'<Trabajador NumeroDocumento="1020304050" PrimerNombre="ANA" PrimerApellido="RUIZ"/><ComprobanteTotal>2500000.00</ComprobanteTotal></NominaIndividual>')
    n = datos_xml(nomina)
    assert (n["tipo"], n["folio"], n["nit_emisor"], n["nombre_receptor"], n["total"]) == ("Nómina Individual", "NI70", EMPRESA, "ANA RUIZ", Decimal("2500000.00"))


def test_lista_de_cufe_se_organiza_con_los_datos_del_xml(tmp_path, capsys):
    venta, compra, malo = "a1" * 48, "b2" * 48, "c3" * 48
    lista = tmp_path / "cufes.txt"
    lista.write_text(f"Lista de CUFE\n1) {venta}\n{compra.upper()}, {malo}\n{venta}\n", encoding="utf-8")
    docs = leer_token(lista)
    assert [d.cufe.lower() for d in docs] == [venta, compra, malo] and not docs[0].completo
    fuente = DianFalsa({venta: _factura(venta, "FE164", "2026-09-20", EMPRESA, "800100200", "357000").encode(),
                        compra.upper(): _adjunto(_factura(compra, "FV10", "2026-09-15", "900111222", EMPRESA, "1190000")),
                        malo: RuntimeError("código 66")})
    destino = tmp_path / "out"
    assert main(["--desde-excel", str(lista), "--destino", str(destino), "--nit", EMPRESA, "--pausa", "0"], cliente=fuente) == 2
    assert (destino / "2026" / "09" / "ventas" / "FE164_800100200.xml").exists()
    assert (destino / "2026" / "09" / "compras" / "FV10_900111222.xml").exists()
    hoja = load_workbook(destino / "indice.xlsx")["Indice"]
    assert hoja.cell(2, 3).value == "FE164" and hoja.cell(2, 9).value == 357000 and hoja.cell(3, 6).value == "EMISOR SAS"
    capsys.readouterr()

    # Segunda corrida: lo que quedó en el índice no se vuelve a pedir y conserva sus datos.
    fuente2 = DianFalsa({malo: RuntimeError("código 66")})
    assert main(["--desde-excel", str(lista), "--destino", str(destino), "--nit", EMPRESA, "--pausa", "0"], cliente=fuente2) == 2
    assert fuente2.pedidos == [malo] and "ya estaban 2" in capsys.readouterr().out
    assert load_workbook(destino / "indice.xlsx")["Indice"].cell(2, 3).value == "FE164"


def test_lista_en_excel_y_sin_nit(tmp_path, capsys):
    wb = Workbook()
    wb.active.append(["CUFE"])
    wb.active.append(["d4" * 48])
    ruta = tmp_path / "lista.xlsx"
    wb.save(ruta)
    assert [d.cufe for d in leer_token(ruta)] == ["d4" * 48]
    assert main(["--desde-excel", str(ruta), "--destino", str(tmp_path / "out")], cliente=DianFalsa({})) == 1
    assert "--nit" in capsys.readouterr().err

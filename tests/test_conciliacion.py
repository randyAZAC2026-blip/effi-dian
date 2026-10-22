import json
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacion import controles as c
from conciliacion.datos import leer_periodo
from conciliacion.dian import verificar_en_dian
from conciliacion.excel import escribir_excel
from conciliacion.modelos import Datos, Parametros

FIXTURE = json.loads((Path(__file__).parent / "fixtures" / "periodo_sep2026.json").read_text(encoding="utf-8"))
P = Parametros(hoy=date(2026, 10, 10))


@pytest.fixture
def datos() -> Datos:
    d = Datos(desde=date(2026, 9, 1), hasta=date(2026, 9, 30))
    for k, v in FIXTURE.items():
        if not k.startswith("_"):
            setattr(d, k, v)
    return d


def docs(hallazgos):
    return sorted(h.documento for h in hallazgos)


# ---------------- Ventas ----------------
def test_factura_sin_cufe_despues_de_la_gracia(datos):
    # FE-161 lleva 37 días sin CUFE; FE-162 es de ayer (en gracia); FE-160 tiene CUFE.
    assert docs(c.ventas_sin_cufe(datos, P)) == ["FE-161"]


def test_prefijo_electronico_con_ebill_cero(datos):
    hs = c.ventas_prefijo_sin_ebill(datos, P)
    assert docs(hs) == ["FE-164"]          # POS-30/31 no son electrónicas
    assert "FE" in hs[0].detalle and hs[0].valor == Decimal("357000")


def test_huecos_del_consecutivo_cuentan_las_anuladas(datos):
    # FE-160..166: faltan 163 y 165 (166 está anulada, pero existe).
    assert docs(c.ventas_huecos_consecutivo(datos, P)) == ["FE-163", "FE-165"]


def test_nota_electronica_sin_cude(datos):
    assert docs(c.notas_venta_sin_cude(datos, P)) == ["NC-8"]


# ---------------- Compras ----------------
def test_recibido_en_el_buzon_y_no_registrado(datos):
    # PROV-B no está en factura_c; PROV-E tampoco (reciente, igual se reporta); PROV-D está rechazado: no se exige registrarlo.
    assert docs(c.compras_no_registradas(datos, P)) == ["N-2", "PP-77"]


def test_factura_de_compra_sin_cufe_o_fuera_del_buzon(datos):
    assert docs(c.compras_sin_cufe_o_fuera_del_buzon(datos, P)) == ["MANUAL-3", "OTRO-MES"]


def test_diferencia_de_valor_mayor_a_un_peso(datos):
    hs = c.compras_diferencia_valor(datos, P)
    assert docs(hs) == ["T-5"]              # 85.000 en Effi vs 80.000 en el XML; T-5 bis cuadra
    assert "5,000.00" in hs[0].detalle


def test_radian_pendiente_despues_de_10_dias(datos):
    assert docs(c.compras_radian_pendiente(datos, P)) == ["T-5"]   # PROV-E (estado 2) tiene 5 días


def test_rechazada_en_radian_pero_registrada(datos):
    assert docs(c.compras_rechazada_registrada(datos, P)) == ["R-1"]


def test_cufe_de_proveedor_repetido(datos):
    hs = c.compras_cufe_repetido(datos, P)
    assert docs(hs) == ["T-5", "T-5 bis"]
    assert "2 facturas" in hs[0].detalle


def test_documento_soporte_sin_cuds(datos):
    assert docs(c.soporte_sin_cuds(datos, P)) == ["DSE-27"]


# ---------------- Nómina ----------------
def test_nomina_sin_cune(datos):
    hs = c.nomina_sin_cune(datos, P)
    assert len(hs) == 1 and hs[0].tercero == "ANA PEREZ"


def test_nomina_anulada_sin_nota_de_ajuste(datos):
    hs = c.nomina_anulada_sin_ajuste(datos, P)
    assert [h.tercero for h in hs] == ["SOFIA RIOS"]   # PEDRO DIAZ sí tiene cufe_ajuste


# ---------------- Negativos y conjunto ----------------
def test_sin_datos_no_hay_hallazgos():
    vacio = Datos(desde=date(2026, 9, 1), hasta=date(2026, 9, 30))
    assert c.conciliar(vacio, P, {"ventas", "compras", "nomina"}) == []


def test_solo_corren_los_modulos_pedidos(datos):
    assert {h.modulo for h in c.conciliar(datos, P, {"nomina"})} == {"Nómina"}
    assert len(c.conciliar(datos, P, {"ventas", "compras", "nomina"})) == 17   # 5 ventas, 10 compras, 2 nómina


def test_la_gracia_es_configurable(datos):
    assert docs(c.ventas_sin_cufe(datos, Parametros(hoy=date(2026, 10, 10), dias_gracia_cufe=0))) == ["FE-161", "FE-162"]


# ---------------- Excel ----------------
def test_excel_con_resumen_y_una_hoja_por_control(datos, tmp_path):
    hs = c.conciliar(datos, P, {"ventas", "compras", "nomina"})
    ruta = escribir_excel(hs, tmp_path / "salida" / "conciliacion.xlsx", "Septiembre 2026")
    wb = load_workbook(ruta)
    assert wb.sheetnames[0] == "Resumen"
    assert len(wb.sheetnames) == 1 + len({h.control for h in hs})
    resumen = wb["Resumen"]
    filas = list(resumen.iter_rows(min_row=2, max_col=5, values_only=True))
    assert filas[0][2] == "Alta"                                   # las altas primero
    assert sum(f[3] for f in filas if f[3]) == len(hs)
    hoja = wb["Recibida y no registrada"]
    assert hoja.auto_filter.ref and hoja["A1"].value == "Módulo"
    compras = wb["Compra sin CUFE del buzón"]
    assert compras.max_row == 3
    buzon = wb["Evento RADIAN pendiente"]
    assert buzon["I2"].hyperlink is None
    sin_registro = [r for r in hoja.iter_rows(min_row=2, values_only=True)]
    assert {r[2] for r in sin_registro} == {"PP-77", "N-2"}


def test_excel_sin_hallazgos(tmp_path):
    wb = load_workbook(escribir_excel([], tmp_path / "vacio.xlsx"))
    assert wb.sheetnames == ["Resumen"] and wb["Resumen"]["A2"].value == "Sin hallazgos en el periodo."


def test_enlace_al_pdf(datos, tmp_path):
    datos.facturas_c = [f for f in datos.facturas_c if f["cufe_proveedor"] != "PROV-A"]
    hs = c.compras_no_registradas(datos, P)
    wb = load_workbook(escribir_excel(hs, tmp_path / "x.xlsx"))
    ws = wb["Recibida y no registrada"]
    fila = next(r for r in ws.iter_rows(min_row=2) if r[2].value == "FV-10")
    assert fila[8].hyperlink.target == "https://effi.com.co/pdf/1"


# ---------------- Lectura y DIAN ----------------
class ClienteFalso:
    def __init__(self):
        self.llamadas = []

    def listar_todo(self, recurso, **filtros):
        self.llamadas.append((recurso, filtros))
        return iter([{"id": "1", "recurso": recurso}])


def test_lee_vigentes_y_anuladas_y_el_buzon_en_tramos_de_31_dias():
    cli = ClienteFalso()
    d = leer_periodo(cli, date(2026, 7, 1), date(2026, 9, 30), {"ventas", "compras", "nomina"})
    buzon = [f for r, f in cli.llamadas if r == "documento_ebill_proveedor"]
    assert len(buzon) == 3 and all((f["hasta"] - f["desde"]).days < 31 for f in buzon)
    assert [f.get("vigente") for r, f in cli.llamadas if r == "factura_v"] == [1, 0]
    assert [x.get("_anulada", False) for x in d.facturas_v] == [False, True]
    assert d.nominas[1]["_anulada"] is True


def test_solo_lee_los_modulos_pedidos():
    cli = ClienteFalso()
    leer_periodo(cli, date(2026, 9, 1), date(2026, 9, 30), {"nomina"})
    assert {r for r, _ in cli.llamadas} == {"nomina"}


class DianFalsa:
    def get_status(self, codigo):
        if codigo == "c160":
            return {"IsValid": "false", "StatusDescription": "Documento con errores"}
        if codigo == "cune71":
            raise RuntimeError("Código 66: no existe en este ambiente")
        return {"IsValid": "true"}


def test_verificacion_en_la_dian(datos):
    hs = verificar_en_dian(datos, DianFalsa())
    assert sorted(h.detalle.split(" · ")[0] for h in hs) == ["GetStatus: Documento con errores", "GetStatus: Error al consultar: Código 66: no existe en este ambiente"]
    assert all(h.severidad == "Alta" for h in hs)

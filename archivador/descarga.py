"""Descarga en lote de los XML con GetXmlByDocumentKey: idempotente, valida el CUFE de cada XML y organiza
<destino>/<año>/<mes>/<clase>/<número>_<NIT>.xml. Escribe indice.xlsx y, si se pide, un ZIP para el Tablero."""
from __future__ import annotations

import re
import time
import zipfile
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Callable, Iterable, Protocol

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .token import Documento, clase, fecha, nit, valor


class FuenteXml(Protocol):
    def get_xml(self, track_id: str) -> bytes: ...


@dataclass
class Resultado:
    doc: Documento
    clase: str
    ruta: Path | None
    estado: str          # descargado · ya_estaba · error
    detalle: str = ""


def _seguro(t: str) -> str:
    return re.sub(r"[^\w.-]+", "_", t).strip("_") or "documento"


def ruta_de(destino: Path, d: Documento, cls: str) -> Path:
    anio, mes = (f"{d.fecha:%Y}", f"{d.fecha:%m}") if d.fecha else ("sin_fecha", "00")
    # El otro: en compras, el proveedor que emite; en ventas, nómina y documento soporte la empresa emite y el tercero recibe.
    tercero = d.nit_emisor if cls == "compras" else d.nit_receptor
    return destino / anio / mes / cls / f"{_seguro(d.numero or d.cufe[:16])}_{tercero or 'sin_nit'}.xml"


def trae_cufe(xml: bytes, cufe: str) -> bool:
    """¿El XML es de ese documento? Su cbc:UUID (o el de la factura dentro de un AttachedDocument) debe ser el CUFE."""
    c = cufe.strip().lower()
    try:
        raiz = etree.fromstring(xml)
    except etree.XMLSyntaxError:
        return False
    uuids = [(x.text or "").strip().lower() for x in raiz.iter() if isinstance(x.tag, str) and etree.QName(x).localname == "UUID"]
    if c in uuids:
        return True
    # AttachedDocument: la factura va como texto dentro de cbc:Description.
    return any(c in (x.text or "").lower() for x in raiz.iter() if isinstance(x.tag, str) and etree.QName(x).localname == "Description")


def _hijos(el: etree._Element, *ruta: str) -> list[etree._Element]:
    """Los descendientes que siguen la ruta de nombres locales (sin importar el namespace)."""
    nivel = [el]
    for nombre in ruta:
        nivel = [x for n in nivel for x in n if isinstance(x.tag, str) and etree.QName(x).localname == nombre]
    return nivel


def _t(el: etree._Element, *ruta: str) -> str:
    x = _hijos(el, *ruta)
    return (x[0].text or "").strip() if x else ""


def _nit_parte(el: etree._Element, parte: str) -> tuple[str, str]:
    p = _hijos(el, parte, "Party")
    if not p:
        return "", ""
    p = p[0]
    return (nit(_t(p, "PartyTaxScheme", "CompanyID") or _t(p, "PartyLegalEntity", "CompanyID")),
            _t(p, "PartyTaxScheme", "RegistrationName") or _t(p, "PartyLegalEntity", "RegistrationName") or _t(p, "PartyName", "Name"))


def datos_xml(xml: bytes) -> dict:
    """Número, fecha, tipo, emisor, receptor y total del XML (factura, notas, documento soporte, nómina o
    AttachedDocument con la factura adentro). Lo que no encuentre queda vacío."""
    raiz = etree.fromstring(xml)
    tipo = etree.QName(raiz).localname
    if tipo == "AttachedDocument":
        for d in raiz.iter():
            texto = (d.text or "").strip() if isinstance(d.tag, str) and etree.QName(d).localname == "Description" else ""
            if texto.startswith("<"):
                try:
                    return datos_xml(texto.encode())
                except etree.XMLSyntaxError:
                    continue
        emisor, nombre_e = nit(_t(raiz, "SenderParty", "PartyTaxScheme", "CompanyID")), _t(raiz, "SenderParty", "PartyTaxScheme", "RegistrationName")
        receptor, nombre_r = nit(_t(raiz, "ReceiverParty", "PartyTaxScheme", "CompanyID")), _t(raiz, "ReceiverParty", "PartyTaxScheme", "RegistrationName")
        return {"tipo": "Factura electrónica", "folio": _t(raiz, "ParentDocumentID") or _t(raiz, "ID"), "fecha": fecha(_t(raiz, "IssueDate")),
                "nit_emisor": emisor, "nombre_emisor": nombre_e, "nit_receptor": receptor, "nombre_receptor": nombre_r}
    if "Nomina" in tipo:
        atr = lambda ruta, a: (_hijos(raiz, ruta)[0].get(a) or "").strip() if _hijos(raiz, ruta) else ""  # noqa: E731
        return {"tipo": "Nómina Individual", "folio": atr("NumeroSecuenciaXML", "Numero"), "fecha": fecha(atr("InformacionGeneral", "FechaGen")),
                "nit_emisor": nit(atr("Empleador", "NIT")), "nombre_emisor": atr("Empleador", "RazonSocial"),
                "nit_receptor": nit(atr("Trabajador", "NumeroDocumento")),
                "nombre_receptor": " ".join(x for x in (atr("Trabajador", "PrimerNombre"), atr("Trabajador", "PrimerApellido")) if x),
                "total": valor(_t(raiz, "ComprobanteTotal"))}
    emisor, nombre_e = _nit_parte(raiz, "AccountingSupplierParty")
    receptor, nombre_r = _nit_parte(raiz, "AccountingCustomerParty")
    if tipo == "Invoice":
        tipo = "Documento soporte" if _t(raiz, "InvoiceTypeCode") == "05" else "Factura electrónica"
    else:
        tipo = {"CreditNote": "Nota crédito", "DebitNote": "Nota débito"}.get(tipo, tipo)
    return {"tipo": tipo, "folio": _t(raiz, "ID"), "fecha": fecha(_t(raiz, "IssueDate")), "nit_emisor": emisor, "nombre_emisor": nombre_e,
            "nit_receptor": receptor, "nombre_receptor": nombre_r, "total": valor(_t(raiz, "LegalMonetaryTotal", "PayableAmount"))}


def completar(d: Documento, xml: bytes) -> Documento:
    """Llena con el XML lo que no trae un documento de una lista de CUFE."""
    try:
        datos = datos_xml(xml)
    except etree.XMLSyntaxError:
        return d
    return replace(d, **{k: v for k, v in datos.items() if v and not getattr(d, k)})


def _previos(destino: Path) -> dict[str, Path]:
    """CUFE → archivo de las corridas anteriores (del indice.xlsx), para no volver a pedir los de una lista."""
    indice = destino / "indice.xlsx"
    if not indice.is_file():
        return {}
    try:
        filas = list(load_workbook(indice, read_only=True).worksheets[0].iter_rows(values_only=True))
    except Exception:  # un índice dañado solo hace que se vuelvan a pedir
        return {}
    if not filas or "CUFE" not in filas[0] or "Archivo" not in filas[0]:
        return {}
    c, a = filas[0].index("CUFE"), filas[0].index("Archivo")
    return {str(f[c]).lower(): Path(str(f[a])) for f in filas[1:] if f[c] and f[a]}


def descargar(docs: Iterable[Documento], fuente: FuenteXml, destino: str | Path, empresa: str, pausa: float = 1.0,
              dormir: Callable[[float], None] = time.sleep, avisar: Callable[[int, int, Resultado], None] | None = None,
              guardar: Callable[[list[Resultado]], None] | None = None, guardar_cada: int = 100) -> list[Resultado]:
    """guardar recibe lo hecho cada guardar_cada descargas y al terminar, también si se corta (Ctrl+C), para retomar."""
    destino = Path(destino)
    docs = list(docs)
    previos = _previos(destino) if any(not d.completo for d in docs) else {}
    out: list[Resultado] = []
    pedidos = 0
    try:
        for i, d in enumerate(docs, start=1):
            previo = previos.get(d.cufe.lower()) if not d.completo else ruta_de(destino, d, clase(d, empresa))
            if previo and previo.exists():
                d = d if d.completo else completar(d, previo.read_bytes())
                r = Resultado(d, clase(d, empresa), previo, "ya_estaba")
            else:
                if pedidos and pausa:
                    dormir(pausa)
                pedidos += 1
                try:
                    xml = fuente.get_xml(d.cufe)
                    if not trae_cufe(xml, d.cufe):
                        raise ValueError("el XML que entregó la DIAN no trae ese CUFE")
                    d = d if d.completo else completar(d, xml)
                    cls = clase(d, empresa)
                    ruta = ruta_de(destino, d, cls)
                    ruta.parent.mkdir(parents=True, exist_ok=True)
                    ruta.write_bytes(xml)
                    r = Resultado(d, cls, ruta, "descargado")
                except Exception as e:  # un documento que falla no detiene el lote
                    r = Resultado(d, clase(d, empresa), None, "error", str(e))
            out.append(r)
            if avisar:
                avisar(i, len(docs), r)
            if guardar and r.estado != "ya_estaba" and pedidos % guardar_cada == 0:
                guardar(out)
    finally:
        if guardar:
            guardar(out)
    return out


def escribir_indice(resultados: list[Resultado], ruta: str | Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Indice"
    ws.append(["Clase", "Tipo", "Número", "Fecha", "NIT emisor", "Emisor", "NIT receptor", "Receptor", "Total", "CUFE", "Archivo", "Estado"])
    for r in resultados:
        d = r.doc
        ws.append([r.clase, d.tipo, d.numero, d.fecha, d.nit_emisor, d.nombre_emisor, d.nit_receptor, d.nombre_receptor, float(d.total), d.cufe,
                   str(r.ruta or ""), {"descargado": "Descargado", "ya_estaba": "Ya estaba", "error": "Error"}[r.estado]])
    fallas = wb.create_sheet("No descargados")
    fallas.append(["Número", "Fecha", "Emisor", "Total", "CUFE", "Motivo"])
    for r in resultados:
        if r.estado == "error":
            fallas.append([r.doc.numero, r.doc.fecha, r.doc.nombre_emisor, float(r.doc.total), r.doc.cufe, r.detalle])
    for hoja in (ws, fallas):
        for c in hoja[1]:
            c.font = Font(bold=True)
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
    destino = Path(ruta)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def escribir_zip(resultados: list[Resultado], ruta: str | Path, base: str | Path) -> Path:
    """ZIP con los XML del lote (descargados o que ya estaban), listo para soltar en el Tablero."""
    destino = Path(ruta)
    destino.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as z:
        for r in resultados:
            if r.ruta and r.ruta.exists():
                z.write(r.ruta, r.ruta.relative_to(base).as_posix())
    return destino

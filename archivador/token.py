"""Lectura del Excel del token de la DIAN (Documentos → Recibidos/Enviados → Exportar), con los mismos nombres de
columna que reconoce la plataforma (frontend/src/utils/tokenDian.js). Los eventos (acuses, aceptación) se omiten."""
from __future__ import annotations

import csv
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

COLUMNAS = {
    "tipo": ["tipo de documento", "tipo documento", "tipo"],
    "cufe": ["cufe/cude", "cufe", "cude", "uuid"],
    "folio": ["folio", "numero", "nro documento", "documento"],
    "prefijo": ["prefijo", "serie"],
    "fecha": ["fecha emision", "fecha de emision", "fecha"],
    "nit_emisor": ["nit emisor", "emisor nit"],
    "nombre_emisor": ["nombre emisor", "razon social emisor"],
    "nit_receptor": ["nit receptor", "receptor nit"],
    "nombre_receptor": ["nombre receptor", "razon social receptor"],
    "total": ["total", "total documento", "valor total"],
}


def plano(s: object) -> str:
    t = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", t).strip()


def es_evento(tipo: str) -> bool:
    return bool(re.search(r"application ?response|evento|acuse", plano(tipo)))


def nit(v: object) -> str:
    return re.sub(r"\D", "", str(v or "").split("-")[0])


def fecha(v: object) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = str(v or "").strip()
    for patron in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(t[:10], patron).date()
        except ValueError:
            continue
    return None


def valor(v: object) -> Decimal:
    t = str(v or "").strip().replace("$", "").replace(" ", "")
    if t.count(",") and t.count("."):
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif t.count(","):
        t = t.replace(",", ".")
    try:
        return Decimal(t or 0)
    except InvalidOperation:
        return Decimal(0)


@dataclass(frozen=True)
class Documento:
    cufe: str
    tipo: str
    folio: str
    prefijo: str
    fecha: date | None
    nit_emisor: str
    nombre_emisor: str
    nit_receptor: str
    nombre_receptor: str
    total: Decimal

    @property
    def completo(self) -> bool:
        """¿Viene del token (con fecha y número)? Los de una lista de CUFE se completan con su XML."""
        return bool(self.fecha or self.folio)

    @property
    def numero(self) -> str:
        return f"{self.prefijo}{self.folio}" if self.prefijo and not str(self.folio).startswith(self.prefijo) else str(self.folio)


def _indices(fila: list[object]) -> dict[str, int] | None:
    titulos = [plano(c) for c in fila]
    idx: dict[str, int] = {}
    for campo, alias in COLUMNAS.items():
        for a in alias:
            if a in titulos:
                idx[campo] = titulos.index(a)
                break
    return idx if "cufe" in idx and "total" in idx else None


def leer_filas(filas: list[list[object]]) -> list[Documento]:
    for i, fila in enumerate(filas[:30]):
        idx = _indices(fila)
        if idx:
            break
    else:
        raise ValueError("No parece el Excel del token de la DIAN: falta la fila con «CUFE/CUDE» y «Total».")
    docs, vistos = [], set()
    for f in filas[i + 1:]:
        v = lambda c: f[idx[c]] if c in idx and idx[c] < len(f) else ""  # noqa: E731
        cufe = str(v("cufe") or "").strip()
        if not cufe or es_evento(str(v("tipo"))) or cufe.lower() in vistos:
            continue
        vistos.add(cufe.lower())
        docs.append(Documento(cufe=cufe, tipo=str(v("tipo") or "").strip(), folio=str(v("folio") or "").strip(), prefijo=str(v("prefijo") or "").strip(),
                              fecha=fecha(v("fecha")), nit_emisor=nit(v("nit_emisor")), nombre_emisor=str(v("nombre_emisor") or "").strip(),
                              nit_receptor=nit(v("nit_receptor")), nombre_receptor=str(v("nombre_receptor") or "").strip(), total=valor(v("total"))))
    return docs


CODIGO = re.compile(r"\b[0-9a-fA-F]{64,}\b")   # CUFE, CUDE y CUNE: SHA-384 en hexadecimal (96); SHA-256 en los viejos


def codigos(texto: str) -> list[str]:
    """Los CUFE de un texto cualquiera (uno por línea, separados por comas, con títulos o numeración), sin repetir."""
    vistos: dict[str, str] = {}
    for c in CODIGO.findall(texto):
        vistos.setdefault(c.lower(), c)
    return list(vistos.values())


def documento_suelto(cufe: str) -> Documento:
    return Documento(cufe=cufe, tipo="", folio="", prefijo="", fecha=None, nit_emisor="", nombre_emisor="", nit_receptor="", nombre_receptor="", total=Decimal(0))


def _filas_csv(ruta: Path) -> list[list[object]]:
    with ruta.open(encoding="utf-8-sig", newline="") as fh:
        muestra = fh.read(4096); fh.seek(0)
        return [list(r) for r in csv.reader(fh, delimiter=";" if muestra.count(";") > muestra.count(",") else ",")]


def leer_token(ruta: str | Path) -> list[Documento]:
    """El Excel o CSV del token. Si no lo es, sirve una lista de CUFE (.txt, .csv o .xlsx): los datos salen del XML."""
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".txt":
        docs = [documento_suelto(c) for c in codigos(ruta.read_text(encoding="utf-8-sig", errors="replace"))]
        if not docs:
            raise ValueError(f"{ruta.name} no trae ningún CUFE (códigos de 96 caracteres).")
        return docs
    if ruta.suffix.lower() == ".csv":
        filas = _filas_csv(ruta)
        try:
            return leer_filas(filas)
        except ValueError as e:
            error = e
    else:
        libro = load_workbook(ruta, read_only=True, data_only=True)
        filas = []
        for hoja in libro.worksheets:
            hoja_filas = [list(r) for r in hoja.iter_rows(values_only=True)]
            try:
                return leer_filas(hoja_filas)
            except ValueError:
                filas += hoja_filas
        error = ValueError("Ninguna hoja tiene la fila de títulos con «CUFE/CUDE» y «Total».")
    sueltos = codigos("\n".join(str(c) for f in filas for c in f if c is not None))
    if not sueltos:
        raise ValueError(f"{error} Tampoco trae una lista de CUFE.")
    return [documento_suelto(c) for c in sueltos]


def nit_empresa(docs: list[Documento]) -> str:
    """La empresa del token: el NIT que más aparece (es emisor o receptor en todos los documentos)."""
    c = Counter(n for d in docs for n in {d.nit_emisor, d.nit_receptor} if n)
    return c.most_common(1)[0][0] if c else ""


def clase(d: Documento, empresa: str) -> str:
    t = plano(d.tipo)
    if "nomina" in t:
        return "nomina"
    if "soporte" in t:
        return "documento_soporte"
    return "ventas" if d.nit_emisor == empresa else "compras"


def leer_solo(ruta: str | Path) -> set[str]:
    """Los CUFE del CSV de faltantes que exporta la plataforma (columna «cufe»)."""
    with Path(ruta).open(encoding="utf-8-sig", newline="") as fh:
        muestra = fh.read(4096); fh.seek(0)
        lector = csv.DictReader(fh, delimiter=";" if muestra.count(";") > muestra.count(",") else ",")
        campo = next((c for c in (lector.fieldnames or []) if plano(c) in ("cufe", "cufe/cude", "cude")), None)
        if not campo:
            raise ValueError("El CSV de faltantes no tiene la columna «cufe».")
        return {str(f[campo]).strip().lower() for f in lector if str(f.get(campo) or "").strip()}

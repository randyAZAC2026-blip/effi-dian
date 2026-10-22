"""GetStatus de una lista de CUFE: una línea por documento, conteo por regla y un Excel que se puede retomar."""
from __future__ import annotations

import re
import time
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .client import DianError

COLUMNAS = ["CUFE", "Válido", "Código", "Descripción", "Reglas", "Notificaciones", "Error"]
GUARDAR_CADA = 100


@dataclass
class Fila:
    cufe: str
    valido: bool = False
    codigo: str = ""
    descripcion: str = ""
    reglas: list[str] = field(default_factory=list)
    notificaciones: list[str] = field(default_factory=list)
    error: str = ""


def reglas(mensajes: list[str]) -> list[str]:
    return [m.group(1) for m in (re.search(r"Regla:\s*(\w+)", x) for x in mensajes) if m]


def leer_previas(ruta: Path) -> dict[str, Fila]:
    """Las filas de una corrida anterior (para retomar). Un archivo dañado o ajeno se ignora."""
    if not ruta.is_file():
        return {}
    try:
        filas = list(load_workbook(ruta, read_only=True).worksheets[0].iter_rows(values_only=True))
    except Exception:
        return {}
    if not filas or list(filas[0][:len(COLUMNAS)]) != COLUMNAS:
        return {}
    out = {}
    for f in filas[1:]:
        f = [("" if v is None else str(v)) for v in f] + [""] * len(COLUMNAS)
        if f[0]:
            out[f[0].lower()] = Fila(cufe=f[0], valido=f[1] == "Sí", codigo=f[2], descripcion=f[3], reglas=[r for r in f[4].split(", ") if r],
                                     notificaciones=[n for n in f[5].split("\n") if n], error=f[6])
    return out


def escribir(filas: list[Fila], ruta: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Estado DIAN"
    ws.append(COLUMNAS)
    for f in filas:
        ws.append([f.cufe, "Sí" if f.valido else "No", f.codigo, f.descripcion, ", ".join(f.reglas), "\n".join(f.notificaciones), f.error])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta


def consultar(cliente, codigos: list[str], salida: Path | None, pausa: float = 1.0, dormir: Callable[[float], None] = time.sleep,
              imprimir: Callable[[str], None] = print) -> list[Fila]:
    previas = leer_previas(salida) if salida else {}
    filas: list[Fila] = []
    pedidos = 0
    n = len(codigos)
    try:
        for i, codigo in enumerate(codigos, start=1):
            previa = previas.get(codigo.lower())
            if previa and not previa.error:
                filas.append(previa)
                continue
            if pedidos and pausa:
                dormir(pausa)
            pedidos += 1
            fila = Fila(cufe=codigo)
            try:
                e = cliente.get_status(codigo)
                fila.valido, fila.codigo, fila.descripcion = e.is_valid, e.status_code, e.status_description
                fila.notificaciones, fila.reglas = e.errores, reglas(e.errores)
            except DianError as error:  # uno que falla no detiene la lista
                fila.error = str(error)
            filas.append(fila)
            linea = f"[{i}/{n}] {codigo[:16]}… "
            if fila.error:
                imprimir(f"{linea}ERROR · {fila.error}")
            elif fila.valido:
                imprimir(f"{linea}VÁLIDO · {fila.codigo}")
            else:
                imprimir(f"{linea}NO VÁLIDO · {fila.codigo} · {fila.descripcion}")
                for m in fila.notificaciones:
                    imprimir(f"  - {m}")
            if salida and pedidos % GUARDAR_CADA == 0:
                _guardar(filas, previas, salida, imprimir)
    except KeyboardInterrupt:
        imprimir("Interrumpido: se guarda lo consultado; vuelve a correr el mismo comando para seguir.")
        raise
    finally:
        if salida:
            _guardar(filas, previas, salida, imprimir)
    return filas


def _guardar(filas: list[Fila], previas: dict[str, Fila], salida: Path, imprimir: Callable[[str], None]) -> None:
    """Lo de esta corrida más lo de corridas anteriores que no está en esta lista."""
    vistos = {f.cufe.lower() for f in filas}
    try:
        escribir(filas + [p for c, p in previas.items() if c not in vistos], salida)
    except PermissionError:  # en Windows, el archivo abierto en Excel no se puede reescribir
        imprimir(f"No se pudo guardar {salida}: ciérralo en Excel; se intenta de nuevo más adelante.")


def _n(x: int) -> str:
    return f"{x:,}".replace(",", ".")


def resumen(filas: list[Fila]) -> list[str]:
    c = Counter("error" if f.error else "valido" if f.valido else "no_valido" for f in filas)
    lineas = [f"Válidos {_n(c['valido'])} de {_n(len(filas))} · no válidos {_n(c['no_valido'])} · con error {_n(c['error'])}"]
    por_regla = Counter(r for f in filas for r in set(f.reglas))
    for regla, cuantos in por_regla.most_common():
        ejemplo = next(m for f in filas for m in f.notificaciones if re.search(rf"Regla:\s*{regla}\b", m))
        lineas.append(f"  {regla}: {_n(cuantos)} documentos · {ejemplo.split('Notificación:')[-1].strip()}")
    return lineas

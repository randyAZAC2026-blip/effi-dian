"""Excel de hallazgos: hoja «Resumen» (conteo y valor por control y severidad) y una hoja por control."""
from __future__ import annotations

from collections import defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .modelos import ORDEN_SEVERIDAD, Hallazgo

COLUMNAS = ("Módulo", "Severidad", "Documento", "Tercero", "NIT", "Valor", "Fecha", "Detalle", "Enlace")
ANCHOS = (11, 10, 18, 32, 14, 16, 12, 70, 12)
COLOR = {"Alta": "F8D7DA", "Media": "FFF3CD", "Baja": "E2E3E5"}
NEGRITA = Font(bold=True)
INVALIDOS = str.maketrans({c: " " for c in "[]:*?/\\"})


def _nombre_hoja(nombre: str, usados: set[str]) -> str:
    base = nombre.translate(INVALIDOS)[:31].strip()
    n, candidato = 2, base
    while candidato in usados:
        sufijo = f" ({n})"
        candidato = base[: 31 - len(sufijo)] + sufijo
        n += 1
    usados.add(candidato)
    return candidato


def _encabezado(ws, columnas, anchos) -> None:
    ws.append(list(columnas))
    for i, ancho in enumerate(anchos, start=1):
        ws.cell(row=1, column=i).font = NEGRITA
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"


def escribir_excel(hallazgos: list[Hallazgo], ruta: str | Path, titulo: str = "") -> Path:
    wb = Workbook()
    resumen = wb.active
    resumen.title = "Resumen"
    _encabezado(resumen, ("Control", "Módulo", "Severidad", "Cantidad", "Valor"), (46, 11, 10, 10, 18))
    if titulo:
        resumen["G1"] = titulo
        resumen["G1"].font = Font(bold=True, size=12)

    por_control: dict[str, list[Hallazgo]] = defaultdict(list)
    for h in hallazgos:
        por_control[h.control].append(h)
    orden = sorted(por_control, key=lambda c: (ORDEN_SEVERIDAD.get(por_control[c][0].severidad, 9), por_control[c][0].modulo, c))
    for control in orden:
        hs = por_control[control]
        total = sum((h.valor for h in hs), Decimal(0))
        resumen.append([control, hs[0].modulo, hs[0].severidad, len(hs), float(total)])
        resumen.cell(row=resumen.max_row, column=3).fill = PatternFill("solid", fgColor=COLOR.get(hs[0].severidad, "FFFFFF"))
        resumen.cell(row=resumen.max_row, column=5).number_format = "#,##0.00"
    if not hallazgos:
        resumen.append(["Sin hallazgos en el periodo."])

    usados = {"Resumen"}
    for control in orden:
        ws = wb.create_sheet(_nombre_hoja(control, usados))
        _encabezado(ws, COLUMNAS, ANCHOS)
        for h in hallazgos_ordenados(por_control[control]):
            ws.append([h.modulo, h.severidad, h.documento, h.tercero, h.nit, float(h.valor), h.fecha, h.detalle, "PDF" if h.enlace else ""])
            fila = ws.max_row
            ws.cell(row=fila, column=6).number_format = "#,##0.00"
            ws.cell(row=fila, column=7).number_format = "yyyy-mm-dd"
            if h.enlace:
                ws.cell(row=fila, column=9).hyperlink = h.enlace
                ws.cell(row=fila, column=9).style = "Hyperlink"
        ws.auto_filter.ref = ws.dimensions

    destino = Path(ruta)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino



def hallazgos_ordenados(hs: list[Hallazgo]) -> list[Hallazgo]:
    return sorted(hs, key=lambda h: (h.fecha or date.min, h.documento))
